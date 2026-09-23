# -*- coding: utf-8 -*-
"""
每一頁都要能在離線、沒有金鑰的情況下打開。

頁面拆成檔案後，最常見的壞法是某頁少 import 一個東西，只有點到那頁才會爆。
"""
from __future__ import annotations

from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parent.parent
PAGES = sorted(p.relative_to(ROOT).as_posix() for p in (ROOT / "views").glob("*.py")
               if p.name != "__init__.py")


@pytest.fixture(autouse=True)
def _offline(monkeypatch, tmp_path):
    monkeypatch.setenv("LLM_PROVIDER", "none")
    import sys
    sys.path.insert(0, str(ROOT / "src"))
    import planner_settings
    monkeypatch.setattr(planner_settings, "DEFAULT_DB", tmp_path / "ps.db")


def test_home_page_runs():
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    assert not at.exception, at.exception


@pytest.mark.parametrize("page", PAGES)
def test_every_page_runs(page):
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page(page).run()
    assert not at.exception, (page, at.exception)


# ---------------------------------------------------------------------------
# Task 4：企劃確認交期表單、兩個匯出按鈕、供應商月度績效
# ---------------------------------------------------------------------------
def _first_needs_review_row():
    """
    直接用底層函式（不透過 Streamlit）找一張目前 needs_human_review 的單，
    跟頁面上看到的資料是同一份（沒有覆寫、沒有確認的乾淨狀態）。
    比起在畫面上用 head(12) 亂猜哪張單會被展開，這樣才能保證測試找到的
    po_no／sched_line／batch_key／email_id 是真的存在、而且跟表單 key
    用的是同一組值。

    batch_key 多數是空字串（這個排程行沒有撞批）；只有供應商提議拆批、
    或規則 4 保守退路撞批的列才有值——見 pipeline._resolve_record_batches。
    """
    import pipeline as _pipeline
    from domain import coalesce_batch_key, coalesce_sched_line
    result = _pipeline.run(use_llm=False)
    tcfg = _pipeline.load_config()["triage"]
    actions = _pipeline.retriage(result.get("all", result["actions"]), {}, tcfg)
    row = actions[actions["needs_human_review"]].iloc[0]
    return (row["po_no"], coalesce_sched_line(row.get("sched_line")),
            coalesce_batch_key(row.get("batch_key")), row["email_id"])


def test_confirmation_form_clears_review_flag_and_logs(tmp_path):
    """
    守住：企劃在畫面上填表單送出後，rerun 一次分級就要立即改用確認日期
    （不再是需人工確認），而且確認紀錄真的寫進資料庫——不能表單顯示
    「已登錄」，資料庫卻是空的。
    """
    from datetime import date

    po_no, sched_line, batch_key, email_id = _first_needs_review_row()

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/actions.py").run()

    # 目標單不一定落在預設篩選（P1、P2）裡：展開所有優先級＋只看需人工確認，
    # 才能保證這張單的展開區（含表單）真的被渲染出來。
    at.multiselect(key="action_priority_pick").set_value(["P1", "P2", "P3", "待查"])
    at.checkbox(key="action_only_review").set_value(True)
    at.run()

    at.get_by_key(f"confirm_date_{po_no}_{sched_line}_{batch_key}_{email_id}").set_value(date(2026, 10, 1))
    at.get_by_key(f"confirm_user_{po_no}_{sched_line}_{batch_key}_{email_id}").set_value("王小明")
    at.get_by_key(f"confirm_submit_{po_no}_{sched_line}_{batch_key}_{email_id}").click()
    at.run()
    assert not at.exception, at.exception

    import planner_settings as ps
    log = ps.confirmation_log()
    assert len(log) == 1
    assert log[0]["po_no"] == po_no and log[0]["confirmed_by"] == "王小明"
    assert log[0]["confirmed_date"] == "2026-10-01"

    # 確認後這張單不再需人工確認：因為 only_review 勾著，它現在應該從
    # 「只看需人工確認」的篩選結果裡消失（表單／按鈕都不該再出現）。
    at2 = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at2.switch_page("views/actions.py").run()
    at2.multiselect(key="action_priority_pick").set_value(["P1", "P2", "P3", "待查"])
    at2.checkbox(key="action_only_review").set_value(True)
    at2.run()
    assert not at2.exception, at2.exception
    remaining_keys = {b.key for b in at2.button}
    assert f"confirm_submit_{po_no}_{sched_line}_{batch_key}_{email_id}" not in remaining_keys

    # 表單消失只是畫面上的側面證據；真正要守住的是分級本身確實變了——
    # 直接用底層函式重算一次（跟畫面用的是同一份邏輯），檢查這張單的
    # needs_human_review 旗標與第一條理由，而不是只看「按鈕還在不在」
    # 這種容易因為改版面就巧合通過的弱驗證。
    import pipeline as _pipeline
    from domain import coalesce_batch_key, coalesce_sched_line
    result2 = _pipeline.run(use_llm=False)
    tcfg2 = _pipeline.load_config()["triage"]
    recomputed = _pipeline.retriage(
        result2.get("all", result2["actions"]), {}, tcfg2,
        confirmations=ps.load_confirmations())
    # 用 (po_no, sched_line, batch_key) 當鍵：分批交貨的單同一個
    # (po_no, sched_line) 可能有兩列（供應商提議拆批），只用前兩者當
    # 索引會撞上重複索引，.loc[key] 撈到的可能是另一批。
    key = (po_no, sched_line, batch_key)
    sched_col = recomputed["sched_line"].map(coalesce_sched_line)
    batch_col = recomputed["batch_key"].map(coalesce_batch_key)
    recomputed_keys = set(zip(recomputed["po_no"], sched_col, batch_col))
    if key in recomputed_keys:
        r = recomputed.set_index(["po_no", sched_col, batch_col]).loc[key]
        assert not r["needs_human_review"]
        assert "王小明" in r["reasons"][0] and "確認交期 2026-10-01" in r["reasons"][0]


def test_action_page_has_both_download_buttons_and_excel_is_readable():
    """
    守住：CSV 匯出（操作指引已提到，文字不能改）與新的 Excel 匯出都要在，
    且 Excel 真的讀得回來（欄位對得上，不是壞檔）。
    """
    import io

    from openpyxl import load_workbook

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/actions.py").run()
    assert not at.exception, at.exception

    labels = {b.label for b in at.download_button}
    assert "⬇️ 匯出行動清單 CSV" in labels
    assert "⬇️ 匯出明日追料清單（Excel）" in labels

    import exports as _exports
    import pipeline as _pipeline

    result = _pipeline.run(use_llm=False)
    tcfg = _pipeline.load_config()["triage"]
    actions = _pipeline.retriage(result.get("all", result["actions"]), {}, tcfg)
    wb = load_workbook(io.BytesIO(_exports.followup_workbook(actions, str(result["as_of"]))))
    assert "追料清單" in wb.sheetnames and "說明" in wb.sheetnames


def test_supplier_page_runs_and_month_select_has_options():
    """守住：月度績效的月份選單要有選項可選，不能是空的下拉。"""
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/suppliers.py").run()
    assert not at.exception, at.exception

    month_boxes = [s for s in at.selectbox if s.label == "月份"]
    assert month_boxes and len(month_boxes[0].options) > 0

    dl_labels = {b.label for b in at.download_button}
    assert "⬇️ 匯出供應商月度績效（Excel）" in dl_labels


# ---------------------------------------------------------------------------
# 強模型審查回饋（Plan 2 Task 1-4 review）：I1-I7、MINOR 的畫面驗證
# ---------------------------------------------------------------------------
def test_today_confirmed_expander_lists_confirmation_with_status(tmp_path):
    """
    I2：登錄過的確認交期要能在「今日已確認」找到，即使那張單現在已經
    不在清單的目前篩選條件裡。I1：這裡的「狀態」欄要標「生效中」——
    email_id 對得上目前這張單最新一封信，這筆確認還算數。
    """
    po_no, sched_line, batch_key, email_id = _first_needs_review_row()
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/actions.py").run()
    at.multiselect(key="action_priority_pick").set_value(["P1", "P2", "P3", "待查"])
    at.checkbox(key="action_only_review").set_value(True)
    at.run()
    at.get_by_key(f"confirm_date_{po_no}_{sched_line}_{batch_key}_{email_id}").set_value(
        __import__("datetime").date(2026, 10, 1))
    at.get_by_key(f"confirm_user_{po_no}_{sched_line}_{batch_key}_{email_id}").set_value("王小明")
    at.get_by_key(f"confirm_submit_{po_no}_{sched_line}_{batch_key}_{email_id}").click()
    at.run()

    at2 = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at2.switch_page("views/actions.py").run()
    assert not at2.exception, at2.exception
    today_confirmed = [ex for ex in at2.expander if ex.label.startswith("📌 今日已確認")]
    assert today_confirmed and today_confirmed[0].label == "📌 今日已確認（1）"
    tables = [df.value for df in at2.dataframe]
    conf_table = next(df for df in tables if "採購單號" in df.columns and "狀態" in df.columns)
    row = conf_table.set_index("採購單號").loc[po_no]
    assert row["狀態"] == "生效中"
    assert row["確認人"] == "王小明"


def test_discoverability_hint_and_uncapped_review_list():
    """
    I4：需人工確認的數量旁要有「怎麼登錄確認」的提示；勾選「只看需人工
    確認」後，清單不能被寫死的 12 筆上限擋住看不到、也填不到後面幾張的
    確認表單。目前合成資料有 14 張需人工確認的單（> 12），足以驗證上限
    真的被拿掉。
    """
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/actions.py").run()
    caps = [c.value for c in at.caption]
    assert any("只看需人工確認" in c and "登錄確認" in c for c in caps)

    at.multiselect(key="action_priority_pick").set_value(["P1", "P2", "P3", "待查"])
    at.checkbox(key="action_only_review").set_value(True)
    at.run()
    assert not at.exception, at.exception
    review_expanders = [ex for ex in at.expander if not ex.label.startswith("📌")]
    assert len(review_expanders) > 12


def test_confirmation_uses_toast_not_lost_success():
    """
    I5：st.success 接著 st.rerun() 在同一次互動裡會被蓋掉，企劃看不到
    「已存成功」；改用 st.toast，跨這次 rerun 還留著。
    """
    po_no, sched_line, batch_key, email_id = _first_needs_review_row()
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/actions.py").run()
    at.multiselect(key="action_priority_pick").set_value(["P1", "P2", "P3", "待查"])
    at.checkbox(key="action_only_review").set_value(True)
    at.run()
    at.get_by_key(f"confirm_date_{po_no}_{sched_line}_{batch_key}_{email_id}").set_value(
        __import__("datetime").date(2026, 10, 1))
    at.get_by_key(f"confirm_user_{po_no}_{sched_line}_{batch_key}_{email_id}").set_value("王小明")
    at.get_by_key(f"confirm_submit_{po_no}_{sched_line}_{batch_key}_{email_id}").click()
    at.run()
    assert not at.exception, at.exception
    assert any("已登錄確認" in t.value for t in at.toast)


def test_receiving_override_uses_toast_not_lost_success():
    """I5：收貨處理天數頁的調整表單也要用 st.toast，理由同上。"""
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/receiving.py").run()
    at.selectbox[0].set_value(at.selectbox[0].options[0])
    at.number_input[0].set_value(5)
    at.text_input[0].set_value("測試原因")
    at.text_input[1].set_value("測試員")
    at.button[0].click()
    at.run()
    assert not at.exception, at.exception
    assert any("已更新" in t.value for t in at.toast)


def test_unmatched_row_guard_present_in_actions_view():
    """
    I3：目前合成資料沒有 matched=False 又 needs_human_review 的單可以
    直接在畫面上點，因此改用原始碼檢查守住這個防呆：對不到主檔的列
    不能顯示確認表單，只能顯示「請先確認單號」的提示。這條防呆一旦被
    誤刪，pipeline.retriage() 對這種列本來就不會套用確認（見
    tests/test_confirmations.py），畫面卻還讓企劃填一個永遠不會生效的
    表單，比不能填更容易誤導人。
    """
    src = (ROOT / "views" / "actions.py").read_text(encoding="utf-8")
    assert 'row.get("matched"' in src
    assert "信中的採購單號對不到系統" in src


def test_supplier_monthly_table_has_overdue_column_and_dash_for_nan():
    """
    I6／I7：月度績效表要有「逾期未收（筆）」欄；樣本不足的月份，天數欄
    要顯示「—」而不是空白或 NaN 字樣，企劃才知道「這裡沒有數字」不是
    「忘記填」。
    """
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    at.switch_page("views/suppliers.py").run()
    assert not at.exception, at.exception
    tables = [df.value for df in at.dataframe]
    monthly = next(df for df in tables if "逾期未收（筆）" in df.columns)
    assert "供應商名稱" in monthly.columns and monthly.columns[0] == "供應商名稱"
    small = monthly[monthly["樣本"] == "樣本不足"]
    if not small.empty:
        assert (small["延遲時中位數(天)"] == "—").all()
        assert (small["P80 延遲(天)"] == "—").all()
